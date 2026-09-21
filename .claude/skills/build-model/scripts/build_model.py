"""Create a model for a company and bring it to the repo's standard.

One command does the whole job:

  1. copy `base.xlsx` to companies/<TICKER>.xlsx (dots in tickers become underscores)
  2. add the company to portfolio.xlsx `Main` and link the ticker to the model
  3. set the header row to span exactly the company's public life
  4. link each quarter to its earnings release, falling back to the 10-Q/10-K
  5. put the 10-Q/10-K in row 1 above each quarter
  6. pin the actuals/forecast arrows to Q326 and 2026

    python build_model.py CRWV
    python build_model.py FSLR --name "First Solar, Inc."
    python build_model.py FSLR --dry-run

Set SEC_EDGAR_USER_AGENT to a contact email first.
"""
import argparse, os, re, shutil, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
LINKS = os.path.abspath(os.path.join(HERE, "..", "..", "refresh-links", "scripts"))
sys.path.insert(0, LINKS)
import refresh_links as R          # noqa: E402
import xlsx_safe                   # noqa: E402

TEMPLATE = "base.xlsx"
INDEX = "portfolio.xlsx"


def sheet_ticker(ticker):
    """Ticker as the index sheet writes it; the file uses underscores for dots."""
    return ticker.strip().upper()


def file_stem(ticker):
    return sheet_ticker(ticker).replace(".", "_")


def company_profile(ticker):
    """Name and industry straight from EDGAR, so the index row is consistent
    with how the company actually reports."""
    tmap = R.ticker_map()
    t = sheet_ticker(ticker)
    if t not in tmap:
        return None
    cik, name = tmap[t]
    sub = R.json.loads(R.fetch("https://data.sec.gov/submissions/CIK%010d.json" % cik,
                               cap=20000000))
    return {"cik": cik, "name": sub.get("name") or name,
            "industry": sub.get("sicDescription") or None}


def add_index_row(ticker, profile, dry_run=False):
    """Add the company to Main, or link an existing row. Returns a status string.

    The link lives on the Ticker cell (column C); Company (column B) is plain
    text and carries no ticker.
    """
    import openpyxl
    from copy import copy
    if not os.path.exists(INDEX):
        return "no %s to update" % INDEX
    wb = openpyxl.load_workbook(INDEX)
    ws = wb["Main"]
    t = sheet_ticker(ticker)
    target = R.model_path(file_stem(ticker))

    last = 2
    existing = None
    for r in range(3, ws.max_row + 1):
        v = ws.cell(row=r, column=3).value
        if v:
            last = r
            if str(v).strip().upper() == t:
                existing = r
    style_from = None
    for r in range(3, ws.max_row + 1):
        if ws.cell(row=r, column=3).hyperlink:
            style_from = (ws.cell(row=r, column=2), ws.cell(row=r, column=3))
            break

    row = existing or last + 1
    if existing and ws.cell(row=row, column=3).hyperlink:
        return "already on Main at row %d" % row
    if dry_run:
        return "would %s Main row %d" % ("link" if existing else "add", row)

    if not existing:
        ws.cell(row=row, column=2, value=profile["name"] if profile else t)
        ws.cell(row=row, column=3, value=t)
        if profile and profile.get("industry"):
            ws.cell(row=row, column=4, value=profile["industry"])
    cell = ws.cell(row=row, column=3)
    cell.hyperlink = target
    if style_from is not None:                      # plain Company, linked Ticker
        ws.cell(row=row, column=2)._style = copy(style_from[0]._style)
        cell._style = copy(style_from[1]._style)
    xlsx_safe.save_workbook(wb, INDEX)
    return "%s Main row %d" % ("linked" if existing else "added", row)


def run(script, *args):
    cmd = [sys.executable, os.path.join(LINKS, script)] + list(args)
    return subprocess.run(cmd, check=False).returncode


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ticker")
    ap.add_argument("--name", help="company name for the index row (default: EDGAR's)")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not R.UA or "@" not in R.UA:
        sys.exit("Set SEC_EDGAR_USER_AGENT to a contact email -- SEC returns 403 without one.")

    t = sheet_ticker(args.ticker)
    stem = file_stem(args.ticker)
    path = R.model_path(stem)

    profile = company_profile(t)
    if profile is None:
        print("! %s has no EDGAR ticker mapping. The model can still be created, but "
              "nothing can be linked." % t)
    elif args.name:
        profile["name"] = args.name
    if profile:
        print("  %s = %s (CIK %d, %s)" % (t, profile["name"], profile["cik"],
                                          profile.get("industry") or "industry unknown"))

    # 1. the model itself
    if os.path.exists(path):
        print("  %s already exists -- leaving it alone" % path)
    elif args.dry_run:
        print("  would create %s from %s" % (path, TEMPLATE))
    else:
        os.makedirs(R.MODEL_DIR, exist_ok=True)
        shutil.copy2(TEMPLATE, path)
        print("  created %s" % path)

    # 2. the index row
    print("  index: %s" % add_index_row(t, profile, args.dry_run))

    if args.dry_run or profile is None:
        print("\n(dry run)" if args.dry_run else "\nNo EDGAR mapping; skipping links.")
        return

    # 3-5. history range, release links, then the filing row above them
    print("\n-- history range and release links --")
    run("refresh_links.py", stem, "--fallback-periodic")
    print("\n-- 10-Q/10-K links in row 1 --")
    run("refresh_links.py", stem, "--filing-row")

    # 6. arrows last: the steps above move columns underneath them
    print("\n-- actuals/forecast arrows --")
    run("place_arrows.py", stem)

    import openpyxl
    ws = openpyxl.load_workbook(path)["Model"]
    q = [c for c in ws[2] if c.value and re.fullmatch(r"Q[1-4](\d{2}|\d{4})",
                                                      str(c.value).strip())]
    rel = sum(1 for c in q if c.hyperlink)
    fil = sum(1 for c in q if ws.cell(row=1, column=c.column).hyperlink)
    print("\n%s: %d quarters (%s..%s), %d release links, %d filing links"
          % (stem, len(q), q[0].value, q[-1].value, rel, fil) if q else
          "\n%s: no quarter headers" % stem)


if __name__ == "__main__":
    main()
